import io
import re
from dataclasses import dataclass
from pathlib import Path

from docx import Document as DocxDocument
from pypdf import PdfReader

from app.chunks.headings import HeadingMarker
from app.documents.constants import SUPPORTED_DOCUMENT_EXTENSIONS

# python-docx's built-in heading style naming - "Heading 1", "Heading 2", ...
# The captured digit becomes the marker's level. Word's 'Title' style has no
# numbered sibling and is handled separately (see _heading_level_for_style).
_DOCX_HEADING_STYLE_RE = re.compile(r"^Heading (\d+)$")


@dataclass(frozen=True)
class ExtractedDocument:
    """Result of extract_document - `text` has the exact same contract
    extract_text used to (pypdf join for .pdf, paragraph join for .docx,
    verbatim UTF-8 decode for .md/.txt). `headings` carries whatever
    *format-native* structural signal this file's extraction could
    recover (tier 2 for .docx/.pdf, tier 4 for .pdf without an outline) -
    always [] for .md/.txt, and possibly [] for .docx/.pdf too if neither
    signal was present. Tiers 1/3 are NOT computed here - see
    chunks/splitting.py's split_document, which tries those itself
    whenever `headings` is empty."""

    text: str
    headings: list[HeadingMarker]


class UnsupportedFileTypeError(Exception):
    """Raised by extract_document for any extension other than .pdf, .docx,
    .md, .txt (case-insensitive)."""


def extract_document(filename: str, data: bytes) -> ExtractedDocument:
    """Extracts plain text - and, where recoverable, format-native heading
    structure - from `data` based on filename's extension. .pdf via pypdf,
    .docx via python-docx, .md/.txt via UTF-8 decode (no structure signal
    for the latter two - see ExtractedDocument)."""
    extension = Path(filename).suffix.lower()
    if extension not in SUPPORTED_DOCUMENT_EXTENSIONS:
        raise UnsupportedFileTypeError(f"Unsupported file type: {filename}")
    if extension == ".pdf":
        # Stub - Task 4 replaces this with real outline/font-size heading
        # detection (_extract_pdf_document). Text extraction is unchanged.
        return ExtractedDocument(text=_extract_pdf_text(data), headings=[])
    if extension == ".docx":
        return _extract_docx_document(data)
    return ExtractedDocument(text=data.decode("utf-8"), headings=[])


def _extract_pdf_text(data: bytes) -> str:
    reader = PdfReader(io.BytesIO(data))
    return "\n".join(page.extract_text() or "" for page in reader.pages)


def _extract_docx_document(data: bytes) -> ExtractedDocument:
    """Joins paragraph text exactly as the old _extract_docx_text did
    ("\\n".join(p.text for p in document.paragraphs)). While joining,
    tracks each paragraph's running character offset and records a
    HeadingMarker whenever its style is one of Word's built-in 'Heading N'
    styles (level = N) or the 'Title' style (level 1)."""
    document = DocxDocument(io.BytesIO(data))
    lines: list[str] = []
    headings: list[HeadingMarker] = []
    offset = 0
    for paragraph in document.paragraphs:
        style_name = paragraph.style.name if paragraph.style is not None else None
        level = _heading_level_for_style(style_name)
        if level is not None:
            headings.append(HeadingMarker(offset=offset, level=level))
        lines.append(paragraph.text)
        offset += len(paragraph.text) + 1  # +1 for the '\n' join separator
    return ExtractedDocument(text="\n".join(lines), headings=headings)


def _heading_level_for_style(style_name: str | None) -> int | None:
    """Maps a python-docx paragraph style name to a heading level, or None
    if the style isn't a heading style. 'Heading N' -> N; Word's built-in
    'Title' style -> 1 (it has no numbered sibling)."""
    if style_name is None:
        return None
    if style_name == "Title":
        return 1
    match = _DOCX_HEADING_STYLE_RE.match(style_name)
    return int(match.group(1)) if match else None
