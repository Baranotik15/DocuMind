import io
from pathlib import Path

import pytest
from docx import Document as DocxDocument

from app.chunks.headings import HeadingMarker
from app.chunks.splitting import split_into_chunks
from app.documents.extraction import UnsupportedFileTypeError, extract_document

FIXTURES_DIR = Path(__file__).parent / "fixtures"

# Must match the text baked into tests/fixtures/sample.pdf by the fixture
# generator (see the commit message / PR description for how it was made).
PDF_FIXTURE_KNOWN_TEXT = "DocuMind sample PDF fixture for extraction tests."


def _build_docx_bytes(paragraphs: list[str], styles: list[str | None] | None = None) -> bytes:
    document = DocxDocument()
    for index, paragraph in enumerate(paragraphs):
        style = styles[index] if styles else None
        if style:
            document.add_paragraph(paragraph, style=style)
        else:
            document.add_paragraph(paragraph)
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


# --- extract_document: text -------------------------------------------------


def test_extract_text_txt_returns_exact_content() -> None:
    content = "Hello, world!\nSecond line.\n"

    result = extract_document("notes.txt", content.encode("utf-8")).text

    assert result == content


def test_extract_text_md_returns_exact_content() -> None:
    content = "# Heading\n\nBody text with **emphasis**.\n"

    result = extract_document("README.md", content.encode("utf-8")).text

    assert result == content


def test_extract_text_is_case_insensitive_on_extension() -> None:
    content = "Uppercase extension still works."

    result = extract_document("NOTES.TXT", content.encode("utf-8")).text

    assert result == content


def test_extract_text_pdf_returns_known_text() -> None:
    data = (FIXTURES_DIR / "sample.pdf").read_bytes()

    result = extract_document("sample.pdf", data).text

    assert PDF_FIXTURE_KNOWN_TEXT in result


def test_extract_text_docx_returns_paragraph_text() -> None:
    data = _build_docx_bytes(["First paragraph.", "Second paragraph."])

    result = extract_document("notes.docx", data).text

    assert "First paragraph." in result
    assert "Second paragraph." in result


def test_extract_text_raises_for_unsupported_extension() -> None:
    with pytest.raises(UnsupportedFileTypeError):
        extract_document("virus.exe", b"whatever bytes")


def test_extract_text_raises_for_extensionless_filename() -> None:
    with pytest.raises(UnsupportedFileTypeError):
        extract_document("no_extension_at_all", b"whatever bytes")


# --- extract_document: DOCX heading styles (tier 2) -------------------------


def test_extract_document_docx_no_heading_styles_returns_no_headings() -> None:
    data = _build_docx_bytes(["First paragraph.", "Second paragraph."])

    result = extract_document("notes.docx", data)

    assert result.headings == []


def test_extract_document_docx_all_normal_style_returns_no_headings() -> None:
    paragraphs = ["Overview", "Body one.", "Details", "Body two."]
    data = _build_docx_bytes(paragraphs, styles=[None, None, None, None])

    result = extract_document("notes.docx", data)

    assert result.headings == []


def test_extract_document_docx_heading_styles_produce_markers() -> None:
    paragraphs = ["Overview", "Body one.", "Details", "Body two."]
    data = _build_docx_bytes(paragraphs, styles=["Heading 1", None, "Heading 2", None])

    result = extract_document("notes.docx", data)

    overview_offset = result.text.index("Overview")
    details_offset = result.text.index("Details")
    assert result.headings == [
        HeadingMarker(offset=overview_offset, level=1),
        HeadingMarker(offset=details_offset, level=2),
    ]


# --- split_into_chunks -----------------------------------------------------
#
# Hard contract: "".join(split_into_chunks(text)) == text EXACTLY, always -
# no trimming, no whitespace normalization, no dropped characters.
#
# Guarantee this implementation provides beyond the bare minimum: every
# returned chunk (not just all-but-the-last) has length <= max_chars. This
# is achievable because a paragraph longer than max_chars is hard-split
# (at sentence boundaries, falling back to a raw character cut) rather than
# ever being allowed to stand as an oversized chunk.


def test_split_into_chunks_short_text_returns_single_chunk() -> None:
    text = "Just a short paragraph, well under the limit."

    result = split_into_chunks(text, max_chars=1500)

    assert result == [text]
    assert "".join(result) == text


def test_split_into_chunks_multi_paragraph_over_limit_splits_on_paragraphs() -> None:
    paragraph = "Lorem ipsum dolor sit amet. " * 3
    text = "\n\n".join([paragraph] * 5)

    result = split_into_chunks(text, max_chars=200)

    assert "".join(result) == text
    assert len(result) > 1
    for chunk in result:
        assert len(chunk) <= 200


def test_split_into_chunks_huge_single_paragraph_hard_splits() -> None:
    sentence = "This is one sentence in a very long paragraph. "
    text = sentence * 50  # no blank lines anywhere in the whole string

    result = split_into_chunks(text, max_chars=200)

    assert "".join(result) == text
    assert len(result) > 1
    for chunk in result:
        assert len(chunk) <= 200


def test_split_into_chunks_empty_string_returns_empty_list() -> None:
    result = split_into_chunks("")

    assert result == []
    assert "".join(result) == ""


def test_split_into_chunks_uses_default_max_chars() -> None:
    text = "x" * 3000

    result = split_into_chunks(text)

    assert "".join(result) == text
    assert len(result) > 1
    for chunk in result:
        assert len(chunk) <= 1500
